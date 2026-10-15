"""04_qtl_region_snps.py -- QTL/BSA interval SNP extraction + gene-region annotation.

Title:    Rice-skills QTL interval SNP extraction (BSA / QTL-seq follow-up)
Description:
    Given a multi-sample resequencing VCF, a GFF3 annotation and a QTL
    interval (from QTL-seq / BSA mapping, e.g. TBtools or QTLseqR output),
    extract the two target parents' genotypes for every SNP in the interval
    and annotate each SNP against the gene model:

      1. interval extraction (chrom + closed [start, end]) from the VCF,
         keeping only the two requested samples in the output VCF;
      2. per-SNP genotype decode: GT, DP, AD, ALT read count and ALT allele
         frequency (AF) for each parent, plus dAF = AF(A) - AF(B) and a
         same/different genotype call (homozygous vs heterozygous difference);
      3. carrier-origin classification: whether each variant is carried by
         parent A only, parent B only, both, or neither (background donors);
      4. gene-region annotation via the GFF3: promoter (TSS-upstream
         --promoter bp, strand aware), 5'UTR / exon / intron / 3'UTR /
         intragenic-other / intergenic. All transcripts of a gene vote and
         labels are merged with "/"; SNPs in several genes list them all;
      5. outputs: filtered interval VCF + one Excel workbook with three
         sheets (all sites / differential sites only / interval gene list
         with per-region SNP counts) + run_info provenance.

    This is the standard "zoom in on the QTL interval" step between BSA
    mapping and candidate-gene work: the differential-sites sheet directly
    feeds manual candidate prioritization (e.g. non-synonymous SNPs in
    exons, promoter variants of annotated genes).

Usage:
    python scripts/qtl/04_qtl_region_snps.py \
        --vcf  filtered.vcf.gz \
        --gff  annotation.gff3.gz \
        --region Chr7:20000000-25000000 \
        --sample-a ParentA --sample-b ParentB \
        --promoter 2000 --outdir results/qtl/04_region_snps
    # self-test with synthetic VCF + GFF (no inputs needed)
    python scripts/qtl/04_qtl_region_snps.py --demo --outdir results/qtl/04_region_snps_demo

Inputs:
    --vcf       multi-sample VCF (plain or .gz)
    --gff       GFF3 annotation (plain or .gz), gene/mRNA/exon/CDS/UTR
    --region    CHROM:START-END (closed interval, 1-based)
    --sample-a / --sample-b   the two parents to compare (must exist in VCF)

Outputs (--outdir):
    {A}_{B}_{chrom}_{start}-{end}.vcf      interval VCF, two samples kept
    region_snps.xlsx                       3 sheets (or .csv per sheet if
                                           openpyxl is unavailable)
    run_info.txt                           parameters + summary statistics

Notes:
    - openpyxl is optional: without it the three tables are written as CSV.
    - The differential-sites sheet also carries per-SNP gene-region labels,
      so exon (exonic) and promoter hits can be filtered directly in Excel.
"""

from __future__ import annotations

import argparse
import csv
import gzip
import sys
import time
from bisect import bisect_right
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from urllib.parse import unquote

# ---------------------------------------------------------------------------
LABEL_ORDER = ["启动子", "5'UTR", "外显子", "内含子", "3'UTR", "基因内其他"]
PRIO = {"5'UTR": 0, "3'UTR": 1, "外显子": 2, "内含子": 3, "基因内其他": 4, "启动子": 5}


def open_maybe_gz(path: Path):
    if str(path).endswith(".gz"):
        return gzip.open(path, "rt", encoding="utf-8", errors="replace")
    return open(path, "r", encoding="utf-8", errors="replace")


# ==================== genotype helpers ====================
def canon_gt(gt: str):
    """Normalise 0|1 == 1/0; return None for missing/partial genotypes."""
    if gt in (".", "./.", ".|.", "./", "/.", None):
        return None
    g = gt.replace("|", "/")
    a = g.split("/")
    if any(x in (".", "") for x in a):
        return None
    return "/".join(sorted(a))


def carrier_state(c):
    if c is None:
        return None
    a = c.split("/")
    if all(x == "0" for x in a):
        return "none"
    return "hom" if all(x != "0" for x in a) else "het"


def ad_info(ad_str: str):
    """Return (ad_string, alt_reads, alt_frequency) from an AD field."""
    if ad_str in (".", "", None):
        return ("NA", "NA", None)
    parts = ad_str.split(",")
    try:
        nums = [int(x) for x in parts]
    except ValueError:
        return (ad_str, "NA", None)
    if len(nums) < 2:
        return (ad_str, 0, None)
    ref_n, alt_n = nums[0], sum(nums[1:])
    tot = ref_n + alt_n
    return (ad_str, alt_n, (alt_n / tot) if tot > 0 else None)


# ==================== GFF annotation ====================
def load_annotation(gff_path: Path, chrom: str, start: int, end: int, pad: int):
    genes, gene_note = {}, {}
    gene_tx = defaultdict(list)
    tx_info, tx_feat = {}, defaultdict(lambda: defaultdict(list))
    with open_maybe_gz(gff_path) as fh:
        for line in fh:
            if not line or line.startswith("#"):
                continue
            p = line.rstrip("\n").split("\t")
            if len(p) < 9 or p[0] != chrom:
                continue
            s, e = int(p[3]), int(p[4])
            if e < start - pad or s > end + pad:
                continue
            attrs = {}
            for kv in p[8].split(";"):
                if "=" in kv:
                    k, v = kv.split("=", 1)
                    attrs[k] = v
            ftype = p[2]
            if ftype == "gene":
                gid = attrs.get("ID", f"gene{len(genes)}")
                genes[gid] = (s, e, p[6])
                gene_note[gid] = unquote(attrs.get("Note", ""))
            elif ftype in ("mRNA", "transcript"):
                gid, tid = attrs.get("Parent"), attrs.get("ID")
                if gid and tid:
                    gene_tx[gid].append(tid)
                    tx_info[tid] = (s, e, p[6], gid)
            elif ftype in ("exon", "CDS", "five_prime_UTR", "three_prime_UTR"):
                tid = attrs.get("Parent")
                if tid:
                    tx_feat[tid][ftype].append((s, e))

    tx_prep = {}
    for tid, (s, e, strand, gid) in tx_info.items():
        prom = (s - pad, s - 1) if strand == "+" else (e + 1, e + pad)
        ex = sorted(tx_feat[tid].get("exon", []))
        introns = []
        for i in range(len(ex) - 1):
            is_, ie_ = ex[i][1] + 1, ex[i + 1][0] - 1
            if is_ <= ie_:
                introns.append((is_, ie_))
        prep = {"prom": prom, "iv": {}, "starts": {}}
        for key, ft in (("five", "five_prime_UTR"), ("three", "three_prime_UTR"),
                        ("cds", "CDS"), ("exon", "exon")):
            lst = sorted(tx_feat[tid].get(ft, []))
            prep["iv"][key] = lst
            prep["starts"][key] = [x[0] for x in lst]
        prep["iv"]["intron"] = introns
        prep["starts"]["intron"] = [x[0] for x in introns]
        tx_prep[tid] = prep
    return genes, gene_note, gene_tx, tx_prep


def in_intervals(iv, starts, pos: int) -> bool:
    i = bisect_right(starts, pos) - 1
    return i >= 0 and iv[i][0] <= pos <= iv[i][1]


def build_gene_ranges(genes, pad):
    ranges = []
    for gid, (s, e, strand) in genes.items():
        lo, hi = (s - pad, e) if strand == "+" else (s, e + pad)
        ranges.append((lo, hi, gid))
    ranges.sort(key=lambda x: x[0])
    return ranges


def annotate_positions(positions, ranges, genes, gene_tx, tx_prep):
    ann_list, hit_list = [], []
    ptr, n, active = 0, len(ranges), []
    for pos in positions:
        while ptr < n and ranges[ptr][0] <= pos:
            active.append(ranges[ptr])
            ptr += 1
        if active:
            active = [g for g in active if g[1] >= pos]
        hits = []
        for lo, hi, gid in active:
            labels, body_any = set(), False
            for tid in gene_tx.get(gid, []):
                prep = tx_prep[tid]
                iv, st = prep["iv"], prep["starts"]
                lab = None
                if in_intervals(iv["five"], st["five"], pos):
                    lab = "5'UTR"
                elif in_intervals(iv["three"], st["three"], pos):
                    lab = "3'UTR"
                elif in_intervals(iv["cds"], st["cds"], pos):
                    lab = "外显子"
                elif in_intervals(iv["exon"], st["exon"], pos):
                    lab = "外显子"
                elif in_intervals(iv["intron"], st["intron"], pos):
                    lab = "内含子"
                if lab:
                    labels.add(lab)
                    body_any = True
                pr = prep["prom"]
                if pr[0] <= pos <= pr[1]:
                    labels.add("启动子")
            gs, ge, strand = genes[gid]
            if gs <= pos <= ge and not body_any:
                labels.add("基因内其他")
            if labels:
                hits.append((gs, gid, [L for L in LABEL_ORDER if L in labels]))
        hits.sort(key=lambda x: x[0])
        ann_list.append("；".join(f"{gid}:{'/'.join(labs)}" for _, gid, labs in hits)
                        if hits else "基因间区")
        hit_list.append(hits)
    return ann_list, hit_list


# ==================== VCF processing ====================
def process_vcf(vcf_path: Path, chrom: str, start: int, end: int,
                sample_a: str, sample_b: str, out_vcf: Path):
    header = []
    rows_all, positions = [], []
    stats = Counter()
    with open_maybe_gz(vcf_path) as fin:
        for line in fin:
            if line.startswith("##"):
                header.append(line.rstrip("\n"))
                continue
            if line.startswith("#CHROM"):
                cols = line.rstrip("\n").split("\t")
                break
        samples = cols[9:]
        if sample_a not in samples or sample_b not in samples:
            raise SystemExit(f"sample(s) not found in VCF. available: {samples}")
        ia, ib = 9 + samples.index(sample_a), 9 + samples.index(sample_b)
        out_cols = cols[:9] + [sample_a, sample_b]
        with open_maybe_gz(vcf_path) as fh2:
            fh2.readline() if False else None  # noqa - stream handled below
        for line in fin:
            if not line or line.startswith("#"):
                continue
            p = line.rstrip("\n").split("\t")
            if p[0] != chrom:
                continue
            pos = int(p[1])
            if pos < start or pos > end:
                continue
            stats["total"] += 1
            ref, alt = p[3], p[4]
            multi = "是" if len(alt.split(",")) > 1 else "否"
            fmt = p[8].split(":")
            rec = {}
            for name, idx in ((sample_a, ia), (sample_b, ib)):
                d = dict(zip(fmt, p[idx].split(":")))
                gt = d.get("GT", ".")
                dp = d.get("DP", "NA")
                try:
                    dp = int(dp)
                except (TypeError, ValueError):
                    dp = "NA"
                ad, alt_n, af = ad_info(d.get("AD", "."))
                rec[name] = (gt, dp, ad, alt_n, af)
            ga, dpa, ada, ana, afa = rec[sample_a]
            gb, dpb, adb, anb, afb = rec[sample_b]
            ca, cb = canon_gt(ga), canon_gt(gb)
            if ca is None or cb is None:
                diff, dtype = "缺失", ""
            elif ca == cb:
                diff, dtype = "相同", ""
            else:
                diff = "不同"
                dtype = "纯合差异" if (ca in ("0/0", "1/1") and cb in ("0/0", "1/1")) else "杂合差异"
            delta = round(afa - afb, 4) if (afa is not None and afb is not None) else "NA"
            sa, sb = carrier_state(ca), carrier_state(cb)
            if sa is None or sb is None:
                origin = "数据缺失"
            elif sa == "none" and sb == "none":
                origin = "两材料均不携带(来自其他样本)"
            elif sb == "none":
                origin = f"仅{sample_a}携带-" + ("纯合" if sa == "hom" else "杂合")
            elif sa == "none":
                origin = f"仅{sample_b}携带-" + ("纯合" if sb == "hom" else "杂合")
            elif sa == "hom" and sb == "hom":
                origin = "两材料均携带(双纯合)"
            else:
                origin = "两材料均携带(含杂合)"
            if diff == "不同":
                stats["diff"] += 1
            rows_all.append([chrom, pos, None, ref, alt, multi,
                             ga, dpa, ada, ana, (round(afa, 4) if afa is not None else "NA"),
                             gb, dpb, adb, anb, (round(afb, 4) if afb is not None else "NA"),
                             delta, diff, dtype, origin])
            positions.append(pos)

    # write interval VCF (two samples)
    with open(out_vcf, "w", encoding="utf-8", newline="") as fo:
        for h in header:
            fo.write(h + "\n")
        fo.write("\t".join(out_cols) + "\n")
        with open_maybe_gz(vcf_path) as fin2:
            copy = False
            for line in fin2:
                if line.startswith("#CHROM"):
                    copy = True
                    continue
                if not copy or not line or line[0] == "#":
                    continue
                p = line.rstrip("\n").split("\t")
                if p[0] != chrom:
                    continue
                pos = int(p[1])
                if pos < start or pos > end:
                    continue
                fo.write("\t".join(p[:9] + [p[ia], p[ib]]) + "\n")
    return rows_all, positions, stats


# ==================== output writers ====================
def write_tables(outdir: Path, rows_all, gene_rows, sample_a, sample_b, chrom):
    header = ["染色体", "位置", "基因区域注释", "参考碱基", "突变碱基", "是否多等位",
              f"{sample_a}_基因型", f"{sample_a}_深度DP", f"{sample_a}_AD(参考,突变)",
              f"{sample_a}_突变reads", f"{sample_a}_突变频率",
              f"{sample_b}_基因型", f"{sample_b}_深度DP", f"{sample_b}_AD(参考,突变)",
              f"{sample_b}_突变reads", f"{sample_b}_突变频率",
              f"ΔAF({sample_a}-{sample_b})", "基因型是否差异", "差异类型", "突变材料归属"]
    gene_header = ["基因ID", "染色体", "起始", "终止", "链", "转录本数", "基因区SNP合计",
                   "启动子SNP", "5'UTR_SNP", "外显子SNP", "内含子SNP", "3'UTR_SNP",
                   "基因内其他SNP", "功能描述(Note)"]
    diff_rows = [r for r in rows_all if r[17] == "不同"]

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        bold = Font(bold=True)
        fill = PatternFill("solid", fgColor="DDEBF7")
        wrap = Alignment(vertical="top", wrap_text=False)

        def sheet(ws, hdr, rows, widths):
            ws.append(hdr)
            for c in range(1, len(hdr) + 1):
                cell = ws.cell(row=1, column=c)
                cell.font = bold
                cell.fill = fill
            for r in rows:
                ws.append(r)
            for i, w in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = w
            ws.freeze_panes = "A2"

        sheet(wb.active, header, rows_all,
              [8, 12, 46, 10, 10, 10, 10, 10, 18, 12, 12, 10, 10, 18, 12, 12, 14, 12, 12, 30])
        wb.active.title = "全部位点"
        sheet(wb.create_sheet("仅差异位点"), header, diff_rows,
              [8, 12, 46, 10, 10, 10, 10, 10, 18, 12, 12, 10, 10, 18, 12, 12, 14, 12, 12, 30])
        sheet(wb.create_sheet(f"{chrom}区间基因列表"), gene_header, gene_rows,
              [16, 8, 12, 12, 6, 9, 13, 12, 12, 12, 12, 12, 14, 60])
        xlsx = outdir / "region_snps.xlsx"
        wb.save(xlsx)
        return str(xlsx), len(diff_rows)
    except ImportError:
        for name, hdr, rows in (("all_sites", header, rows_all),
                                ("diff_sites", header, diff_rows),
                                ("gene_list", gene_header, gene_rows)):
            with open(outdir / f"{name}.csv", "w", newline="", encoding="utf-8-sig") as fh:
                w = csv.writer(fh)
                w.writerow(hdr)
                w.writerows(rows)
        return str(outdir / "all_sites.csv (openpyxl unavailable -> CSV set)"), len(diff_rows)


# ==================== demo ====================
def make_demo(tmp: Path):
    """Synthetic: gene GENE001 (+) with exon/intron/UTR, promoter at 1000-1099;
    interval Chr1:1000-2000; SNPs at 1050 (promoter, genotype-different),
    1150 (exon 1), 1300 (intron), 1600 (intergenic), 2100 (outside)."""
    gff = tmp / "demo.gff3"
    gff.write_text(
        "##gff-version 3\n"
        "Chr1\tdemo\tgene\t1100\t1500\t.\t+\t.\tID=GENE001;Note=demo%20gene\n"
        "Chr1\tdemo\tmRNA\t1100\t1500\t.\t+\t.\tID=GENE001.1;Parent=GENE001\n"
        "Chr1\tdemo\tfive_prime_UTR\t1100\t1120\t.\t+\t.\tParent=GENE001.1\n"
        "Chr1\tdemo\texon\t1100\t1250\t.\t+\t.\tParent=GENE001.1\n"
        "Chr1\tdemo\tCDS\t1100\t1250\t.\t+\t.\tParent=GENE001.1\n"
        "Chr1\tdemo\texon\t1350\t1500\t.\t+\t.\tParent=GENE001.1\n"
        "Chr1\tdemo\tCDS\t1350\t1480\t.\t+\t.\tParent=GENE001.1\n"
        "Chr1\tdemo\tthree_prime_UTR\t1480\t1500\t.\t+\t.\tParent=GENE001.1\n",
        encoding="utf-8")
    vcf = tmp / "demo.vcf"
    lines = [
        "##fileformat=VCFv4.2",
        '##FORMAT=<ID=GT,Number=1,Type=String,Description="Genotype">',
        '##FORMAT=<ID=DP,Number=1,Type=Integer,Description="Depth">',
        '##FORMAT=<ID=AD,Number=R,Type=Integer,Description="Allelic depths">',
        "#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\tFORMAT\tParentA\tParentB\tDonor3",
        # promoter SNP: A het vs B ref, dAF = 0.5 - 0.0 = 0.5
        "Chr1\t1050\t.\tA\tG\t.\t.\t.\tGT:DP:AD\t0/1:20:10,10\t0/0:18:18,0\t0/1:15:7,8",
        # exon SNP: A hom-alt vs B hom-ref
        "Chr1\t1150\t.\tC\tT\t.\t.\t.\tGT:DP:AD\t1/1:22:0,22\t0/0:20:20,0\t0/0:12:12,0",
        # intron SNP: same in both -> not differential
        "Chr1\t1300\t.\tG\tA\t.\t.\t.\tGT:DP:AD\t0/1:19:9,10\t0/1:21:11,10\t0/1:10:5,5",
        # intergenic SNP
        "Chr1\t1600\t.\tT\tC\t.\t.\t.\tGT:DP:AD\t0/0:17:17,0\t0/1:16:8,8\t0/0:9:9,0",
        # outside interval -> must NOT be extracted
        "Chr1\t2100\t.\tA\tT\t.\t.\t.\tGT:DP:AD\t0/1:14:7,7\t0/1:14:7,7\t0/1:8:4,4",
    ]
    vcf.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return vcf, gff


# ==================== main ====================
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--vcf", default=None)
    ap.add_argument("--gff", default=None)
    ap.add_argument("--region", default=None, help="CHROM:START-END (closed)")
    ap.add_argument("--sample-a", dest="sample_a", default=None)
    ap.add_argument("--sample-b", dest="sample_b", default=None)
    ap.add_argument("--promoter", type=int, default=2000, help="promoter length upstream of TSS")
    ap.add_argument("--outdir", default="results/qtl/04_region_snps")
    ap.add_argument("--demo", action="store_true")
    args = ap.parse_args(argv)

    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    t0 = time.time()

    if args.demo:
        vcf, gff = make_demo(outdir)
        chrom, rng = "Chr1", "1000-2000"
        sa, sb = "ParentA", "ParentB"
        promoter = 1000
    else:
        if not (args.vcf and args.gff and args.region and args.sample_a and args.sample_b):
            ap.error("--vcf --gff --region --sample-a --sample-b are required (or --demo)")
        vcf, gff = Path(args.vcf), Path(args.gff)
        chrom, rng = args.region.split(":")
        sa, sb = args.sample_a, args.sample_b
        promoter = args.promoter
    start, end = (int(x) for x in rng.split("-"))

    print(f"[04q] region {chrom}:{start}-{end}; samples {sa} vs {sb}; promoter {promoter} bp")
    genes, gene_note, gene_tx, tx_prep = load_annotation(gff, chrom, start, end, promoter)
    gene_ranges = build_gene_ranges(genes, promoter)
    print(f"[04q] genes near interval: {len(genes)}; transcripts: {len(tx_prep)}")

    out_vcf = outdir / f"{sa}_{sb}_{chrom}_{start}-{end}.vcf"
    rows_all, positions, stats = process_vcf(vcf, chrom, start, end, sa, sb, out_vcf)
    ann_list, hit_list = annotate_positions(positions, gene_ranges, genes, gene_tx, tx_prep)
    gene_snp_n, gene_label_n = Counter(), defaultdict(Counter)
    for row, ann, hits in zip(rows_all, ann_list, hit_list):
        row[2] = ann
        for _, gid, labels in hits:
            gene_snp_n[gid] += 1
            for L in labels:
                gene_label_n[gid][L] += 1

    gene_rows = []
    for gid, (gs, ge, strand) in sorted(genes.items(), key=lambda kv: kv[1][0]):
        overlapping = gs <= end and ge >= start
        if overlapping or gid in gene_snp_n:
            gene_rows.append([gid, chrom, gs, ge, strand, len(gene_tx.get(gid, [])),
                              gene_snp_n.get(gid, 0),
                              gene_label_n[gid].get("启动子", 0),
                              gene_label_n[gid].get("5'UTR", 0),
                              gene_label_n[gid].get("外显子", 0),
                              gene_label_n[gid].get("内含子", 0),
                              gene_label_n[gid].get("3'UTR", 0),
                              gene_label_n[gid].get("基因内其他", 0),
                              gene_note.get(gid, "")])

    table_path, n_diff = write_tables(outdir, rows_all, gene_rows, sa, sb, chrom)
    main_cat = Counter(r[2] if r[2] == "基因间区" else r[2].split(":")[-1].split("；")[0].split("/")[0]
                       for r in rows_all)

    info = [
        f"region: {chrom}:{start}-{end} (closed)",
        f"samples: {sa} vs {sb}; promoter: {promoter} bp",
        f"timestamp: {datetime.now().isoformat(timespec='seconds')}",
        f"SNPs in interval: {stats['total']}; differential: {n_diff}",
        f"genes near interval: {len(genes)}; genes with SNPs: {len(gene_snp_n)}",
        f"outputs: {out_vcf.name}; {Path(table_path).name}",
        f"elapsed: {time.time() - t0:.1f} s",
    ]
    (outdir / "run_info.txt").write_text("\n".join(info) + "\n", encoding="utf-8")
    print("\n".join("[04q] " + s for s in info))

    if args.demo:
        # --- self-test assertions -------------------------------------
        ann_by_pos = {r[1]: r[2] for r in rows_all}
        assert len(rows_all) == 4, f"expected 4 SNPs in interval, got {len(rows_all)}"
        assert "GENE001:启动子" in ann_by_pos[1050], ann_by_pos
        assert "GENE001:5'UTR/外显子" in ann_by_pos[1150] or "外显子" in ann_by_pos[1150], ann_by_pos
        assert "GENE001:内含子" in ann_by_pos[1300], ann_by_pos
        assert ann_by_pos[1600] == "基因间区", ann_by_pos
        r1050 = next(r for r in rows_all if r[1] == 1050)
        assert r1050[17] == "不同" and abs(r1050[16] - 0.5) < 1e-9, r1050
        assert r1050[19].startswith(f"仅{sa}携带-杂合"), r1050
        r1300 = next(r for r in rows_all if r[1] == 1300)
        assert r1300[17] == "相同" and r1300[18] == "", r1300
        assert any(g[0] == "GENE001" and g[6] == 3 for g in gene_rows), gene_rows
        print("[04q] demo self-test PASSED (interval extraction, 4 SNPs; promoter/exon/"
              "intron/intergenic labels; dAF=0.5; gene SNP counts OK)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
